"""Pull plain text out of files so Keeper can compare real content.

Text formats are read directly, PDFs through pypdf, DOCX through python-docx.
A PDF with no text layer (a scan/photo) is flagged ``no_text_layer`` instead of
being silently treated as empty — Keeper says "I couldn't read this one"
rather than guessing.
"""

from __future__ import annotations

import logging
from pathlib import Path

from .scan import MAX_TEXT_CHARS

# pypdf is chatty about imperfect xrefs in the wild; keep Keeper's output clean.
logging.getLogger("pypdf").setLevel(logging.ERROR)

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".yml",
    ".yaml", ".ini", ".cfg", ".toml", ".html", ".htm", ".xml", ".sql",
    ".py", ".js", ".ts", ".sh", ".bat", ".rtf", ".tex", ".srt", ".vtt",
}

# Formats that are never text — never decode them, no matter what the bytes
# look like. (Decoding JPEG/PNG bytes produced garbage "text" that clustered
# random photos together.)
MEDIA_EXTS = {
    # images
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff",
    ".ico", ".psd", ".heic",
    # video / audio
    ".mp4", ".mkv", ".avi", ".mov", ".webm", ".flv", ".m4v", ".mpg",
    ".mpeg", ".mp3", ".wav", ".flac", ".ogg", ".m4a", ".opus",
    # archives & binaries
    ".zip", ".tar", ".gz", ".bz2", ".xz", ".7z", ".rar", ".iso", ".apk",
    ".jar", ".war", ".whl", ".deb", ".rpm", ".exe", ".dll", ".so", ".bin",
    ".pyc", ".o", ".a", ".db", ".sqlite", ".sqlite3",
    # office (old OLE binaries) & fonts
    ".doc", ".xls", ".ppt", ".pps", ".woff", ".woff2", ".ttf", ".otf", ".eot",
}


def extract_text(path: Path) -> tuple[str, str]:
    """Return ``(text, status)``. Status: ok | binary | no_text_layer | error:*."""
    ext = path.suffix.lower()
    try:
        if ext in MEDIA_EXTS:
            return "", "binary"
        if ext == ".pdf":
            return _pdf_text(path)
        if ext == ".docx":
            return _docx_text(path)
        if ext in TEXT_EXTS or ext == "":
            return _plain_text(path)
        # Unknown extension: sniff it. NUL bytes mean binary.
        with path.open("rb") as fh:
            head = fh.read(8192)
        if b"\x00" in head:
            return "", "binary"
        return _plain_text(path)
    except Exception as exc:  # noqa: BLE001 - one bad file must not kill the scan
        return "", f"error:{type(exc).__name__}: {exc}"


def _plain_text(path: Path) -> tuple[str, str]:
    raw = path.open("rb").read(MAX_TEXT_CHARS)
    if b"\x00" in raw:
        return "", "binary"
    text = raw.decode("utf-8", errors="replace")
    if text.count("\ufffd") / max(len(text), 1) > 0.02:
        # Too many undecodable bytes: try a lossless single-byte codec, then
        # judge the result by how much of it is actually printable.
        text = raw.decode("latin-1")
    sample = text[:8192]
    printable = sum(1 for ch in sample if ch.isprintable() or ch in "\n\r\t")
    if sample and printable / len(sample) < 0.85:
        return "", "binary"
    return text, "ok"


def _pdf_text(path: Path) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            chunks.append(page_text)
        if sum(len(c) for c in chunks) >= MAX_TEXT_CHARS:
            break
    text = "\n".join(chunks)[:MAX_TEXT_CHARS]
    if not text.strip():
        # A real PDF, but nothing extractable = scanned image pages.
        return "", "no_text_layer"
    return text, "ok"


def _docx_text(path: Path) -> tuple[str, str]:
    from docx import Document

    doc = Document(str(path))
    chunks: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    text = "\n".join(chunks)[:MAX_TEXT_CHARS]
    if not text.strip():
        return "", "no_text_layer"
    return text, "ok"
